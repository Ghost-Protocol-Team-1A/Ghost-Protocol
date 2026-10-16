"""Tidy tables of a case for Power BI, Tableau or Excel.

One row per thing, one column per field, IDs as keys so the tables join in any BI tool:
evidence -> events (evidence_id), claims -> claim_citations (claim_id) -> evidence, and so on.
Times are UTC. The workbook opens directly in Power BI ("Excel workbook") and Tableau
("Microsoft Excel"); the zip holds the same tables as CSV.
"""

from __future__ import annotations

import csv
import io
import zipfile
from datetime import UTC, datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from app.services.store import CaseData

Table = tuple[list[str], list[list[object]]]


def _dt(value: str | None) -> datetime | None:
    """A UTC datetime without tzinfo, which BI tools read as a date/time column."""
    if not value:
        return None
    parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
    if parsed.tzinfo is not None:
        parsed = parsed.astimezone(UTC)
    return parsed.replace(tzinfo=None)


def tables(data: CaseData) -> dict[str, Table]:
    cid = data.case.id
    names = {e.id: e.source_name for e in data.evidence}
    conflicts_of: dict[str, list[str]] = {}
    for c in data.conflicts:
        for ev in c.event_ids:
            conflicts_of.setdefault(ev, []).append(c.id)
    signed = data.conclusions[-1] if data.conclusions else None
    return {
        "evidence": (
            [
                "case_id",
                "evidence_id",
                "type",
                "source_name",
                "original_filename",
                "sha256",
                "size_bytes",
                "ingested_at_utc",
                "clock_method",
                "clock_offset_seconds",
                "extraction_method",
                "simulated",
            ],
            [
                [
                    cid,
                    e.id,
                    e.type.value,
                    e.source_name,
                    e.original_filename,
                    e.sha256,
                    e.size_bytes,
                    _dt(e.ingested_at),
                    e.clock.method,
                    e.clock.offset_seconds,
                    e.extraction_method.value,
                    e.simulated,
                ]
                for e in data.evidence
            ],
        ),
        "events": (
            [
                "case_id",
                "event_id",
                "evidence_id",
                "source_name",
                "locator",
                "description",
                "source_time_as_recorded",
                "normalized_time_utc",
                "offset_seconds",
                "state",
                "approximate",
                "uncertainty_seconds",
                "extracted_by",
                "conflict_ids",
            ],
            [
                [
                    cid,
                    e.id,
                    e.evidence_id,
                    names.get(e.evidence_id, ""),
                    e.locator,
                    e.description,
                    e.source_time,
                    _dt(e.normalized_time),
                    e.offset_seconds,
                    e.state.value,
                    e.approximate,
                    e.uncertainty_seconds,
                    e.extracted_by or "",
                    ", ".join(conflicts_of.get(e.id, [])),
                ]
                for e in data.events
            ],
        ),
        "claims": (
            ["case_id", "claim_id", "question_id", "statement", "state", "rationale", "created_by"],
            [
                [
                    cid,
                    c.id,
                    c.question_id or "",
                    c.statement,
                    c.state.value,
                    c.rationale,
                    c.created_by or c.origin,
                ]
                for c in data.claims
            ],
        ),
        "claim_citations": (
            ["case_id", "claim_id", "evidence_id", "relationship", "locator"],
            [
                [cid, c.id, link.evidence_id, link.relationship, link.locator or ""]
                for c in data.claims
                for link in c.links
            ],
        ),
        "conflicts": (
            [
                "case_id",
                "conflict_id",
                "rule_id",
                "type",
                "severity",
                "status",
                "evidence_ids",
                "event_ids",
                "explanation",
                "cannot_tell",
            ],
            [
                [
                    cid,
                    c.id,
                    c.rule_id,
                    c.type.value,
                    c.severity.value,
                    c.status.value,
                    ", ".join(c.evidence_ids),
                    ", ".join(c.event_ids),
                    c.explanation,
                    c.uncertain,
                ]
                for c in data.conflicts
            ],
        ),
        "questions": (
            [
                "case_id",
                "question_id",
                "question",
                "finding",
                "answer",
                "signed_by",
                "signed_at_utc",
            ],
            [[cid, q.id, q.text, *_answer(signed, q.id)] for q in data.questions],
        ),
        "audit": (
            ["case_id", "audit_id", "at_utc", "actor", "action", "entity_ref", "details"],
            [
                [cid, a.id, _dt(a.at), a.actor, a.action, a.entity_ref, a.details]
                for a in data.audit
            ],
        ),
    }


def _answer(signed, question_id: str) -> list[object]:
    if signed is None:
        return ["", "", "", None]
    answer = next((a for a in signed.answers if a.question_id == question_id), None)
    if answer is None:
        return ["", "", "", None]
    return [answer.finding.value, answer.text, signed.signed_by, _dt(signed.signed_at)]


ABOUT = [
    ["Ghost Protocol case export"],
    [""],
    ["Case", "{case_id} · {title}"],
    ["Fictional demo data", "{fictional}"],
    ["Times", "All *_utc columns are UTC. source_time_as_recorded keeps the source's own text."],
    ["Keys", "evidence_id, event_id, claim_id, conflict_id and question_id join the sheets."],
    ["Power BI", "Get data > Excel workbook > select every sheet > Load."],
    [
        "Tableau",
        "Connect > Microsoft Excel > drag sheets onto the canvas; relate on the *_id columns.",
    ],
    ["Caution", "An investigative aid. It does not determine responsibility or guilt."],
]


def workbook(data: CaseData) -> bytes:
    wb = Workbook()
    about = wb.active
    about.title = "about"
    for row in ABOUT:
        about.append(
            [
                cell.format(
                    case_id=data.case.id, title=data.case.title, fictional=data.case.fictional
                )
                for cell in row
            ]
        )
    about["A1"].font = Font(bold=True, size=14)
    about.column_dimensions["A"].width = 22
    about.column_dimensions["B"].width = 90
    header_fill = PatternFill("solid", fgColor="B8F54A")
    for name, (columns, rows) in tables(data).items():
        ws = wb.create_sheet(name)
        ws.append(columns)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="0A0B0D")
            cell.fill = header_fill
        for row in rows:
            ws.append(row)
        ws.freeze_panes = "A2"
        for col, column in zip(ws.columns, columns, strict=True):
            width = max([len(column)] + [len(str(c.value or "")) for c in col[1:200]])
            ws.column_dimensions[col[0].column_letter].width = min(max(12, width + 2), 60)
            if column.endswith("_utc"):
                for c in col[1:]:
                    c.number_format = "yyyy-mm-dd hh:mm:ss"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def csv_zip(data: CaseData) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, (columns, rows) in tables(data).items():
            text = io.StringIO()
            writer = csv.writer(text)
            writer.writerow(columns)
            for row in rows:
                writer.writerow(
                    [
                        "" if v is None else v.isoformat() + "Z" if isinstance(v, datetime) else v
                        for v in row
                    ]
                )
            z.writestr(f"{name}.csv", text.getvalue())
    return buf.getvalue()
